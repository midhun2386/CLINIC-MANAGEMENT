"""Export routes — generate Excel, PDF, or Word documents with standardized 8-column output."""

import io
from datetime import datetime

from flask import Blueprint, request, jsonify, g, send_file

from .models import db, Patient, TreatmentEntry
from .auth import token_required, _log

export_bp = Blueprint('export', __name__, url_prefix='/api/export')

# ── Standard 8-column spec (§4 of Blueprint v2) ────────────────────────
EXPORT_COLUMNS = [
    'Name',
    'Date of Birth',
    'Phone Number',
    'Treatment Name',
    'Next Consultation',
    'Payment',
    'Payment Method',
    'Gender',
]


def _gather_data(body: dict) -> list[dict]:
    """
    Build a flat list of rows (one per treatment entry, patient fields repeated).

    Accepts:
        { "export_type": "records" }          — export all treatment entries
        { "export_type": "patients" }         — export all patients (even without entries)
        { "patient_ids": [1, 2, 3] }          — export specific patient(s)
        { "record_ids": [10, 20] }            — export specific treatment entries
        { "search_query": "text" }            — export matching patients & entries
    """
    rows = []

    export_type = body.get('export_type', '').strip().lower()
    patient_ids = body.get('patient_ids', [])
    record_ids = body.get('record_ids', [])
    search_query = body.get('search_query', '').strip()
    date_from = body.get('date_from', '').strip()
    date_to = body.get('date_to', '').strip()

    def _make_row(patient, entry=None):
        return {
            'Name': patient.name if patient else '',
            'Date of Birth': patient.dob.isoformat() if patient and patient.dob else '',
            'Phone Number': patient.phone if patient else '',
            'Treatment Name': entry.treatment_name or '' if entry else '',
            'Next Consultation': entry.next_consultation.isoformat() if entry and entry.next_consultation else '',
            'Payment': str(entry.payment) if entry and entry.payment is not None else '',
            'Payment Method': entry.payment_method or '' if entry else '',
            'Gender': patient.gender if patient else '',
        }

    # 1. Explicit patient IDs
    if patient_ids:
        patients = Patient.query.filter(Patient.id.in_(patient_ids)).all()
        for p in patients:
            if p.treatment_entries:
                for t in p.treatment_entries:
                    rows.append(_make_row(p, t))
            else:
                rows.append(_make_row(p))
        return rows

    # 2. Explicit record IDs or export_type == 'records'
    if export_type == 'records' or record_ids:
        q = TreatmentEntry.query
        if record_ids:
            q = q.filter(TreatmentEntry.id.in_(record_ids))
        if date_from:
            try:
                q = q.filter(TreatmentEntry.date_of_treatment >= datetime.strptime(date_from, '%Y-%m-%d').date())
            except ValueError:
                pass
        if date_to:
            try:
                q = q.filter(TreatmentEntry.date_of_treatment <= datetime.strptime(date_to, '%Y-%m-%d').date())
            except ValueError:
                pass

        entries = q.order_by(TreatmentEntry.date_of_treatment.desc()).all()
        pid_set = {e.patient_id for e in entries}
        p_map = {p.id: p for p in Patient.query.filter(Patient.id.in_(pid_set)).all()} if pid_set else {}

        for e in entries:
            rows.append(_make_row(p_map.get(e.patient_id), e))
        return rows

    # 3. Search query export
    if search_query:
        patients = Patient.query.filter(
            db.or_(
                Patient.name.ilike(f'%{search_query}%'),
                Patient.phone.ilike(f'%{search_query}%'),
            )
        ).all()
        exported_patient_ids = set()
        for p in patients:
            exported_patient_ids.add(p.id)
            if p.treatment_entries:
                for t in p.treatment_entries:
                    rows.append(_make_row(p, t))
            else:
                rows.append(_make_row(p))

        # Also search treatment entries directly
        matching_entries = TreatmentEntry.query.filter(
            db.or_(
                TreatmentEntry.treatment_name.ilike(f'%{search_query}%'),
                TreatmentEntry.medicine.ilike(f'%{search_query}%'),
                TreatmentEntry.remark.ilike(f'%{search_query}%'),
            )
        ).all()
        for e in matching_entries:
            if e.patient_id not in exported_patient_ids:
                pat = Patient.query.get(e.patient_id)
                rows.append(_make_row(pat, e))
        return rows

    # 4. Default: export all patients with their entries
    patients = Patient.query.order_by(Patient.name.asc()).all()
    for p in patients:
        if p.treatment_entries:
            for t in p.treatment_entries:
                rows.append(_make_row(p, t))
        else:
            rows.append(_make_row(p))

    return rows


# ---------------------------------------------------------------------------
# Generators
# ---------------------------------------------------------------------------

def _generate_xlsx(rows: list[dict]) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = 'Clinic Records'

    if not rows:
        ws.append(['No data to export'])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    headers = EXPORT_COLUMNS

    # Style
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    for row_num, row_data in enumerate(rows, 2):
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=row_num, column=col_num, value=row_data.get(header, ''))
            cell.border = thin_border

    # Auto-width
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _generate_pdf(rows: list[dict]) -> io.BytesIO:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=15*mm, rightMargin=15*mm,
                            topMargin=15*mm, bottomMargin=15*mm)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph('Clinic EMR — Export', styles['Title']))
    elements.append(Paragraph(f'Generated: {datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")}',
                              styles['Normal']))
    elements.append(Spacer(1, 10*mm))

    if not rows:
        elements.append(Paragraph('No data to export.', styles['Normal']))
        doc.build(elements)
        buf.seek(0)
        return buf

    headers = EXPORT_COLUMNS
    table_data = [headers]
    cell_style = styles['Normal']
    cell_style.fontSize = 7
    cell_style.leading = 9

    for row in rows:
        table_data.append([
            Paragraph(str(row.get(h, '')), cell_style) for h in headers
        ])

    col_count = len(headers)
    available_width = landscape(A4)[0] - 30*mm
    col_width = available_width / col_count

    table = Table(table_data, colWidths=[col_width]*col_count)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563EB')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('FONTSIZE', (0, 1), (-1, -1), 7),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    elements.append(table)
    doc.build(elements)
    buf.seek(0)
    return buf


def _generate_docx(rows: list[dict]) -> io.BytesIO:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()
    doc.add_heading('Clinic EMR — Export', level=1)
    doc.add_paragraph(f'Generated: {datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")}')

    if not rows:
        doc.add_paragraph('No data to export.')
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf

    headers = EXPORT_COLUMNS
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(9)

    # Data rows
    for row_idx, row_data in enumerate(rows, 1):
        for col_idx, h in enumerate(headers):
            cell = table.rows[row_idx].cells[col_idx]
            cell.text = str(row_data.get(h, ''))
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# Route
# ---------------------------------------------------------------------------

@export_bp.route('', methods=['POST'])
@token_required
def export_data():
    """
    Export data as xlsx, pdf, or docx with standardized 8-column layout.

    Body:
        format: "xlsx" | "pdf" | "docx"
        patient_ids: [int]   — OR —
        record_ids: [int]    — OR —
        search_query: str
    """
    body = request.get_json(silent=True) or {}
    fmt = body.get('format', 'xlsx').lower()

    if fmt not in ('xlsx', 'pdf', 'docx'):
        return jsonify({'error': 'Format must be xlsx, pdf, or docx'}), 400

    rows = _gather_data(body)

    timestamp = datetime.utcnow().strftime('%Y%m%d_%H%M%S')

    if fmt == 'xlsx':
        buf = _generate_xlsx(rows)
        mimetype = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        filename = f'clinic_export_{timestamp}.xlsx'
    elif fmt == 'pdf':
        buf = _generate_pdf(rows)
        mimetype = 'application/pdf'
        filename = f'clinic_export_{timestamp}.pdf'
    else:
        buf = _generate_docx(rows)
        mimetype = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        filename = f'clinic_export_{timestamp}.docx'

    _log(g.current_user.id, 'EXPORT',
         f'format={fmt} rows={len(rows)}', request.remote_addr)

    return send_file(buf, mimetype=mimetype, as_attachment=True, download_name=filename)
